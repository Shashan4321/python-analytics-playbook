import duckdb
import pandas as pd
from openpyxl import load_workbook

from playbook.etl import EtlConfig, incremental_load, write_partitioned_parquet
from playbook.quality import flag_anomalies, reconcile, validate
from playbook.reporting import INR_FORMAT, build_mis_workbook, summary_table


def test_clean_data_passes_validation(sales):
    result = validate(sales)
    assert result.pass_rate == 1.0 and result.quarantined.empty


def test_bad_rows_are_quarantined_not_fatal(tiny):
    bad = tiny.copy()
    bad.loc[1, "quantity"] = 0  # out of range
    bad.loc[3, "region"] = "Mars"  # not a region
    bad.loc[4, "revenue"] = 9999.0  # breaks revenue = qty x price x (1 - discount)
    result = validate(bad)
    assert sorted(result.quarantined["order_id"]) == [2, 4, 5]
    assert result.valid["order_id"].tolist() == [1, 3]
    assert result.pass_rate == 0.4
    assert not result.failures.empty


def test_reconcile_detects_missing_and_amount_drift(tiny):
    ok = reconcile(tiny, tiny.copy(), "order_id", ["revenue"])
    assert ok["passed"].all()

    target = tiny.drop(index=0).copy()
    target.loc[1, "revenue"] += 5
    checks = reconcile(tiny, target, "order_id", ["revenue"]).set_index("check")
    assert not checks.loc["row_count", "passed"]
    assert checks.loc["missing_in_target", "source"] == 1
    assert not checks.loc["sum_revenue", "passed"]


def test_anomaly_flags_spike_only():
    daily = pd.Series([100.0, 104, 98, 101, 99, 103, 500, 97, 102])
    assert flag_anomalies(daily).tolist() == [False] * 6 + [True] + [False] * 2
    assert not flag_anomalies(pd.Series([5.0] * 10)).any()


def test_incremental_load_is_idempotent(tiny, tmp_path):
    cfg = EtlConfig(warehouse=tmp_path / "wh.duckdb")
    first = incremental_load(tiny.iloc[:3], cfg)
    assert (first.inserted, first.updated, first.total_rows) == (3, 0, 3)

    again = incremental_load(tiny.iloc[:3], cfg)
    assert (again.inserted, again.updated, again.total_rows) == (0, 3, 3)

    changed = tiny.iloc[2:].copy()
    changed.loc[2, "revenue"] = 250.0
    last = incremental_load(changed, cfg)
    assert (last.inserted, last.updated, last.total_rows) == (2, 1, 5)
    assert last.watermark == pd.Timestamp("2026-04-06")

    with duckdb.connect(str(cfg.warehouse)) as con:
        assert con.execute("SELECT revenue FROM fact_sales WHERE order_id = 3").fetchone()[0] == 250.0


def test_incremental_load_dedupes_within_batch(tiny, tmp_path):
    batch = pd.concat([tiny, tiny.tail(1).assign(revenue=123.0)])
    result = incremental_load(batch, EtlConfig(warehouse=tmp_path / "wh.duckdb"))
    assert result.total_rows == 5


def test_partitioned_parquet_layout(tiny, tmp_path):
    parts = write_partitioned_parquet(tiny, tmp_path / "lake")
    assert [p.name for p in parts] == ["fiscal_year=2025", "fiscal_year=2026", "fiscal_year=2027"]
    back = pd.read_parquet(tmp_path / "lake")
    assert len(back) == len(tiny)


def test_summary_compares_latest_two_fiscal_years(sales):
    s = summary_table(sales).set_index("kpi")
    assert list(s.columns) == ["FY25", "FY26", "change"]
    assert s.loc["net_revenue", "FY26"] > s.loc["net_revenue", "FY25"]  # generator has growth


def test_mis_workbook_is_formatted(sales, tmp_path):
    path = build_mis_workbook(sales, tmp_path / "MIS.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Monthly", "Category"]
    monthly = wb["Monthly"]
    assert monthly["A1"].value == "month_start" and monthly["A1"].font.bold
    assert monthly["B2"].number_format == INR_FORMAT
    assert monthly.freeze_panes == "B2"
    assert monthly.conditional_formatting
