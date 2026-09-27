"""Formatted Excel MIS report, the kind finance teams expect every month."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from playbook.fiscal import fiscal_year, fy_label
from playbook.kpis import monthly_revenue, net_sales

# Indian digit grouping: 1,23,45,678 (lakh / crore) instead of 12,345,678.
INR_FORMAT = '[>=10000000]"₹"##\\,##\\,##\\,##0;[>=100000]"₹"##\\,##\\,##0;"₹"##,##0'
PCT_FORMAT = "0.0%;[Red]-0.0%"
HEADER_FILL = PatternFill("solid", fgColor="0F2E5C")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def summary_table(df: pd.DataFrame) -> pd.DataFrame:
    """Headline KPIs for the latest fiscal year vs the one before."""
    sales = df.assign(fiscal_year=fiscal_year(df["order_date"]))
    latest = int(sales["fiscal_year"].max())
    rows = []
    for fy in (latest - 1, latest):
        part = sales.loc[sales["fiscal_year"] == fy]
        net = net_sales(part)
        rows.append(
            {
                "fiscal_year": fy_label(fy),
                "net_revenue": net["revenue"].sum(),
                "orders": net["order_id"].nunique(),
                "customers": net["customer_id"].nunique(),
                "avg_order_value": net["revenue"].sum() / max(net["order_id"].nunique(), 1),
                "return_rate": part["returned"].mean(),
            }
        )
    table = pd.DataFrame(rows).set_index("fiscal_year").T
    prev, curr = table.columns
    table["change"] = table[curr] / table[prev] - 1
    return table.rename_axis("kpi").reset_index()


def category_table(df: pd.DataFrame) -> pd.DataFrame:
    sales = net_sales(df).assign(fiscal_year=fy_label(fiscal_year(df["order_date"])))
    return sales.pivot_table(
        index="category", columns="fiscal_year", values="revenue", aggfunc="sum", observed=True
    ).reset_index()


def _style_sheet(ws, money_cols: set[int], pct_cols: set[int]) -> None:
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column in money_cols:
                cell.number_format = INR_FORMAT
            elif cell.column in pct_cols:
                cell.number_format = PCT_FORMAT
    for idx, column in enumerate(ws.columns, start=1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in column)
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 3, 12), 40)
    ws.freeze_panes = "B2"


def build_mis_workbook(df: pd.DataFrame, path: Path) -> Path:
    """Write Summary, Monthly and Category sheets with Indian number formats and a YoY heatmap."""
    summary = summary_table(df)
    monthly = monthly_revenue(df)[["month_start", "revenue", "mom_pct", "yoy_pct", "fytd_revenue"]]
    monthly = monthly.assign(month_start=monthly["month_start"].dt.strftime("%b-%Y"))
    category = category_table(df)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        monthly.to_excel(writer, sheet_name="Monthly", index=False)
        category.to_excel(writer, sheet_name="Category", index=False)

        ws = writer.sheets["Summary"]
        _style_sheet(ws, money_cols=set(), pct_cols={4})
        kpi_formats = {"return_rate": PCT_FORMAT, "orders": "#,##0", "customers": "#,##0"}
        for row in ws.iter_rows(min_row=2):
            fmt = kpi_formats.get(row[0].value, INR_FORMAT)
            for cell in row[1:3]:
                cell.number_format = fmt

        ws = writer.sheets["Monthly"]
        _style_sheet(ws, money_cols={2, 5}, pct_cols={3, 4})
        last = ws.max_row
        ws.conditional_formatting.add(
            f"D2:D{last}",
            ColorScaleRule(
                start_type="num",
                start_value=-0.3,
                start_color="F8696B",
                mid_type="num",
                mid_value=0,
                mid_color="FFFFFF",
                end_type="num",
                end_value=0.3,
                end_color="63BE7B",
            ),
        )

        ws = writer.sheets["Category"]
        _style_sheet(ws, money_cols=set(range(2, ws.max_column + 1)), pct_cols=set())
    return path
